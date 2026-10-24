from pathlib import Path

from app import services
from app.db import connect
from app.schemas import Feedback


def test_saved_briefings_load_for_all_events():
    for eid in ["gfc_2008", "covid_2020", "ru_ua_2022", "legoland_2022", "svb_2023", "tariff_2025"]:
        b = services.load_saved_briefing(eid)
        assert b is not None and b.event.id == eid


def test_feedback_writes_feedback_and_shadow_log():
    b = services.load_saved_briefing("svb_2023")
    services.submit_feedback(Feedback(run_id=b.run_id, doc_id="PB-BANK", label="useful", user_id="u1"), b)
    conn = connect("sandbox")
    assert conn.execute("SELECT label FROM feedback").fetchall()[0][0] == "useful"
    row = conn.execute("SELECT event_id, doc_id, rank, feedback, mode FROM shadow_log").fetchone()
    conn.close()
    assert tuple(row) == ("svb_2023", "PB-BANK", 1, "useful", "backtest")
    conn = connect("live")
    assert conn.execute("SELECT COUNT(*) FROM shadow_log").fetchone()[0] == 0  # 라이브에 섞이지 않음
    conn.close()


def test_priority_aggregates_shadow_log():
    b = services.load_saved_briefing("svb_2023")
    services.submit_feedback(Feedback(run_id=b.run_id, doc_id="PB-BANK", label="not_fit", user_id="u1"), b)
    pr = services.get_priority()
    assert pr["docs"][0]["doc_id"] == "PB-BANK" and pr["docs"][0]["not_fit"] == 1


def test_approval_recorded_in_trace():
    b = services.load_saved_briefing("svb_2023")
    services.record_approval(b, "approved", "u1", "확인")
    trace = services.get_trace(b.trace_id)
    assert trace[-1]["tool"] == "human_approval" and trace[-1]["output_summary"] == "approved"


def test_backtest_results_available():
    res = services.get_backtest_results()
    assert res["sentinel"]["events"] and res["match"] and "svb_2023" in res["exposure"]


def test_live_returns_none_without_news(monkeypatch):
    signals = {"news": {"clusters": []}}
    assert services.run_live(signals) is None


def test_ui_renders_without_exceptions():
    from streamlit.testing.v1 import AppTest

    app = Path(__file__).resolve().parent.parent / "ui" / "streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=120).run()
    assert not at.exception
    assert len(at.tabs) == 7
    at.sidebar.selectbox[0].set_value("svb_2023").run()
    assert not at.exception
    assert any("SVB" in m.value or "실리콘밸리" in m.value for m in at.markdown)


def test_all_ui_tables_serialize_to_arrow(monkeypatch):
    """UI 표가 Arrow로 직렬화되는지 확인 (실패 시 Streamlit이 경고를 내고 자동 변환함)."""
    import pandas as pd
    import pyarrow as pa
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    original, seen, bad = st.dataframe, [], []

    def checked(df, *args, **kwargs):
        if isinstance(df, pd.DataFrame):
            seen.append(list(df.columns))
            try:
                pa.Table.from_pandas(df, preserve_index=False)
            except Exception as e:  # noqa: BLE001
                bad.append((list(df.columns), str(e)[:100]))
        return original(df, *args, **kwargs)

    monkeypatch.setattr(st, "dataframe", checked)
    app = Path(__file__).resolve().parent.parent / "ui" / "streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=120).run()
    at.sidebar.selectbox[0].set_value("svb_2023").run()
    assert seen and not bad, bad


def test_ui_live_mode_renders_map_with_collected_signals(monkeypatch, tmp_path):
    """라이브 모드: 수집기 파일이 있으면 지도·감시 탭이 그 결과로 그려진다."""
    import json

    from streamlit.testing.v1 import AppTest

    path = tmp_path / "signals_latest.json"
    path.write_text(json.dumps({
        "collected_at": "2026-10-18T22:07:00+00:00", "as_of": "2026-10-16", "level": "NONE", "money_alert": False,
        "market": {"active": False, "triggered": [], "zscores": {"VIXCLS": 0.1}}, "flow": None, "fred_error": None,
        "news": {"active": False, "status": "unavailable", "n_queries": 8, "errors": ["pandemic: 429"], "clusters": [],
                 "coverage": {"timespan": "1d", "countries": {"United States": 5}, "by_query": {}, "errors": [],
                              "status": "ok"}},
    }), encoding="utf-8")
    monkeypatch.setattr(services, "LIVE_SIGNALS_PATH", path)
    sig = services.get_signals()
    assert sig["source"] == "collector" and sig["news"]["status"] == "unavailable" and sig["panel"]

    app = Path(__file__).resolve().parent.parent / "ui" / "streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=120).run()
    radio = {r.label: r for r in at.sidebar.radio}
    assert radio["시장"].value == "local"  # 기본은 로컬(대한민국)
    radio["모드"].set_value("라이브").run()
    assert not at.exception
    at.button[0].click().run()  # 감시 탭 '수집 결과 보기'
    assert not at.exception
    assert any("판단 불가" in e.value for e in at.error)
    {r.label: r for r in at.sidebar.radio}["시장"].set_value("global").run()  # 표시 범위만 바뀌고 데이터는 그대로
    assert not at.exception and any("판단 불가" in e.value for e in at.error)


def test_ui_data_input_tab_converts_sample_and_applies_portfolio():
    from streamlit.testing.v1 import AppTest

    app = Path(__file__).resolve().parent.parent / "ui" / "streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=120).run()
    box = next(s for s in at.selectbox if s.label == "또는 예시 파일")
    box.set_value("broker_balance.xlsx").run()
    assert not at.exception
    next(b for b in at.button if b.label == "이 포트폴리오로 분석").click().run()
    assert not at.exception
    assert at.session_state["uploaded_portfolio"].nodes
    at.run()  # 지도 탭은 입력 탭보다 먼저 그려지므로 다음 실행에서 반영
    assert any("업로드" in c.value for c in at.caption)  # 지도 탭이 업로드 포트폴리오를 쓴다
    assert any(r.label == "포트폴리오" for r in at.radio)  # 브리핑 탭에 업로드 가정 분석 선택이 생긴다
    assert at.session_state["uploaded_holdings"][0]["name"]


def test_briefing_excel_export_has_evidence_sheets():
    import io

    from openpyxl import load_workbook

    from app import export

    b = services.load_saved_briefing("svb_2023")
    chain = [{"단계": "A", "원인": "c", "근거": "e", "결과": "r"}]
    wb = load_workbook(io.BytesIO(export.briefing_xlsx(b, chain)))
    assert {"요약", "판단 근거", "문서 적합도", "조항 판정", "노출 경로", "확인 항목"} <= set(wb.sheetnames)
    assert wb["조항 판정"].max_row > 1


def test_ui_document_upload_section_sets_playbooks():
    from streamlit.testing.v1 import AppTest

    app = Path(__file__).resolve().parent.parent / "ui" / "streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=120).run()
    next(m for m in at.multiselect if m.label == "또는 예시 문서").set_value(["PB-BANK.docx"]).run()
    assert not at.exception
    next(b for b in at.button if b.label == "이 문서로 분석").click().run()
    assert not at.exception and at.session_state["uploaded_playbooks"][0].doc_id == "UP-PB-BANK"
    at.run()
    assert any(b.label.startswith("업로드 데이터로 분석") for b in at.button)  # 브리핑 탭에 가정 분석 버튼
