"""대응 문서 입력 검증용 샘플 (data/ingest/doc_samples/): 합성 플레이북을 실제 문서 양식으로 렌더링 + 정답.

- PB-BANK.docx : 워드 (장 제목 + '제n조' 조문 + 담당 부서 문장)
- PB-LIQ.pdf   : PDF (번호 목록 '1.' '2.' + 줄바꿈으로 끊긴 긴 문장) — Chromium으로 인쇄
- truth.json   : 원래 조항 원문·담당·보호 여부
사용: python scripts/make_doc_samples.py
"""

from __future__ import annotations

import asyncio
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from app.config import ROOT_DIR  # noqa: E402
from app.corpus import load_playbooks  # noqa: E402

OUT = ROOT_DIR / "data" / "ingest" / "doc_samples"


def to_docx(pb, path: Path) -> None:
    doc = Document()
    doc.core_properties.title = pb.title
    doc.add_heading(pb.title, level=0)
    doc.add_heading("제1장 총칙 및 대응 절차", level=1)
    for i, c in enumerate(pb.clauses, start=1):
        doc.add_paragraph(f"제{i}조 {c.text} 담당: {c.owner_role}.")
    doc.save(path)


async def to_pdf(pb, path: Path) -> None:
    from playwright.async_api import async_playwright

    items = "".join(f"<p>{i}. {c.text} 이 업무는 {c.owner_role}에서 수행한다.</p>" for i, c in enumerate(pb.clauses, start=1))
    html = f"<html><head><meta charset='utf-8'><style>body{{font-family:'WenQuanYi Zen Hei';width:420px;font-size:13px}}</style></head>" \
           f"<body><h1>{pb.title}</h1><h2>I. 대응 절차</h2>{items}</body></html>"
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
        page = await b.new_page()
        await page.set_content(html)
        await page.pdf(path=str(path), format="A4")
        await b.close()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    pbs = {p.doc_id: p for p in load_playbooks()}
    to_docx(pbs["PB-BANK"], OUT / "PB-BANK.docx")
    asyncio.run(to_pdf(pbs["PB-LIQ"], OUT / "PB-LIQ.pdf"))
    truth = {d: [{"text": c.text, "owner_role": c.owner_role, "protected": c.protected,
                  "event_types": [e.value for e in c.premises.event_types], "channels": [x.value for x in c.premises.channels]}
                 for c in pbs[d].clauses] for d in ("PB-BANK", "PB-LIQ")}
    (OUT / "truth.json").write_text(json.dumps({"synthetic": True, "docs": truth}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"샘플 → {OUT}")


if __name__ == "__main__":
    main()
