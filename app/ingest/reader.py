"""파일 읽기 — xlsx(openpyxl)·csv. 셀 값과 함께 숫자 서식·병합 셀 같은 구조 정보를 남긴다."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

ENCODINGS = ("utf-8-sig", "cp949", "euc-kr", "latin-1")
LARGE_BYTES = 20 * 1024 * 1024  # 이보다 크면 read_only 스트리밍 (병합 셀 정보 생략)


@dataclass
class Sheet:
    name: str
    rows: list[list]  # 전체 행 (값)
    number_formats: dict[int, str] = field(default_factory=dict)  # 열 index → 대표 숫자 서식
    merged: list[str] = field(default_factory=list)  # 병합 범위 (상단 일부)
    encoding: str | None = None


def read_csv_bytes(data: bytes, name: str = "csv") -> Sheet:
    for enc in ENCODINGS:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    return Sheet(name=name, rows=rows, encoding=enc)


def read_xlsx_bytes(data: bytes, merged_rows: int = 30) -> list[Sheet]:
    from openpyxl import load_workbook

    read_only = len(data) > LARGE_BYTES
    wb = load_workbook(io.BytesIO(data), read_only=read_only, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows, formats = [], {}
        for r_idx, row in enumerate(ws.iter_rows()):
            vals = []
            for c_idx, cell in enumerate(row):
                vals.append(cell.value)
                if isinstance(cell.value, (int, float)) and c_idx not in formats and getattr(cell, "number_format", "General") != "General":
                    formats[c_idx] = cell.number_format
            rows.append(vals)
        merged = [] if read_only else [str(m) for m in ws.merged_cells.ranges if m.min_row <= merged_rows]
        out.append(Sheet(name=ws.title, rows=rows, number_formats=formats, merged=merged))
    return out


def read_file(path_or_bytes: Path | bytes, filename: str | None = None) -> list[Sheet]:
    data = path_or_bytes.read_bytes() if isinstance(path_or_bytes, Path) else path_or_bytes
    name = (filename or (path_or_bytes.name if isinstance(path_or_bytes, Path) else "")).lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return read_xlsx_bytes(data)
    return [read_csv_bytes(data, Path(name).stem if name else "csv")]
