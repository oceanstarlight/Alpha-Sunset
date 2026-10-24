"""기관 위기 대응 문서(PDF·DOCX·TXT) → 플레이북(조항) 변환.

1) 라이브러리로 본문을 읽고(pypdf·python-docx) 2) 번호·조문 패턴으로 조항을 나눈다(규칙, 원문 그대로 보존)
3) LLM은 조항마다 속성(전제 사건 유형·채널·지역·기간, 조치, 담당, 법규 보호 여부)만 채운다. 조항 원문은 LLM이 쓰지 않는다.
4) 코드가 검증한다: 담당은 원문에 나온 부서명만, 보호(법규·보험) 판단은 원문 인용 근거가 있을 때만, 열거형 값은 스키마로.
LLM이 없으면 규칙만으로 조항 분해·담당·보호 후보를 채우고 나머지는 비워 둔다 ('사람 확인 필요' 표시).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Literal

from pydantic import BaseModel

from app.schemas import Channel, Clause, Duration, EventType, Playbook, Premises

# 조항 시작 패턴 (원문 그대로 남긴다)
CLAUSE_START = re.compile(
    r"^\s*(?:"
    r"제\s*\d+\s*조(?:의\s*\d+)?"          # 제3조, 제3조의2
    r"|\d+(?:\.\d+)*[.)]\s"                 # 1. 1.2. 3)
    r"|\(\d+\)\s?"                           # (1)
    r"|[①-⑳]"                               # ①
    r"|[가-하][.)]\s"                        # 가.
    r"|[-•▪◦·]\s"                            # 글머리표
    r")"
)
HEADING_HINT = re.compile(r"^(?:제\s*\d+\s*장|[IVX]+\.|#+\s)")
ROLE_RE = re.compile(r"(?:담당\s*[:：]\s*)?([가-힣A-Za-z]{1,12}(?:팀|부서|본부|위원회|센터))")
PROTECT_WORDS = {"regulatory": ["감독", "보고", "법", "규정", "금융위", "금감원", "공시", "인가", "신고"],
                 "insurance": ["보험", "부보", "예금보험", "보상"]}


@dataclass
class RawClause:
    text: str
    section: str = ""
    loc: str = ""  # 페이지·문단 위치


@dataclass
class DocParse:
    title: str
    paragraphs: list[tuple[str, str, str]]  # (텍스트, 스타일, 위치)
    notes: list[str] = field(default_factory=list)


# ---------- 읽기 ----------


def read_document(data: bytes, filename: str) -> DocParse:
    name = filename.lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        from pypdf import PdfReader

        paras = []
        for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
            paras += [(line, "", f"p.{i}") for line in _pdf_lines(page.extract_text() or "")]
        notes = [] if paras else ["PDF에서 글자를 읽지 못함 (스캔 이미지일 수 있음 → OCR 필요)"]
        return DocParse(Path(filename).stem, paras, notes)
    if name.endswith(".docx") or data[:2] == b"PK":
        from docx import Document

        doc = Document(io.BytesIO(data))
        paras = [(p.text, p.style.name if p.style is not None else "", f"문단 {i}") for i, p in enumerate(doc.paragraphs, 1)
                 if p.text.strip()]
        for t_i, table in enumerate(doc.tables, 1):  # 표 안의 조항도 읽는다
            for r_i, row in enumerate(table.rows, 1):
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append((" | ".join(dict.fromkeys(cells)), "Table", f"표{t_i} 행{r_i}"))
        title = doc.core_properties.title or Path(filename).stem
        return DocParse(title, paras)
    text = data.decode("utf-8-sig", errors="replace")
    return DocParse(Path(filename).stem, [(l.rstrip(), "", f"줄 {i}") for i, l in enumerate(text.splitlines(), 1) if l.strip()])


def _pdf_lines(text: str) -> list[str]:
    """PDF 줄을 정리한다. 폭에 걸려 바뀐 줄(가장 긴 줄의 80% 이상, 한글로 끝나고 다음 줄도 한글로 시작)은
    단어 중간에서 끊긴 것으로 보고 공백 없이 잇는다 (예: '점' + '검' → '점검')."""
    lines = [l.rstrip() for l in text.splitlines() if l.strip()]
    full = max((len(l) for l in lines), default=0)
    out: list[str] = []
    carry = False
    for line in lines:
        if carry and out and re.match(r"^[가-힣]", line.lstrip()) and not CLAUSE_START.match(line):
            out[-1] += line.lstrip()
        else:
            out.append(line)
        carry = bool(full) and len(line) >= 0.8 * full and bool(re.search(r"[가-힣]$", line))
    return out


# ---------- 조항 분해 (규칙) ----------


def _is_heading(text: str, style: str) -> bool:
    if style.lower().startswith(("heading", "title", "제목")):
        return True
    if HEADING_HINT.match(text):
        return True
    return False


def split_clauses(parse: DocParse) -> list[RawClause]:
    """번호·조문으로 시작하는 줄에서 새 조항, 이어지는 줄은 앞 조항에 붙인다. 제목은 절 이름으로만 쓴다."""
    clauses: list[RawClause] = []
    section = ""
    for text, style, loc in parse.paragraphs:
        t = re.sub(r"\s+", " ", text).strip()
        if not t:
            continue
        if _is_heading(t, style) and not CLAUSE_START.match(t):
            section = t[:80]
            continue
        if CLAUSE_START.match(t) or "List" in style or style == "Table":
            clauses.append(RawClause(t, section, loc))
        elif clauses and not t.endswith(":") and len(t) > 0:
            clauses[-1].text = f"{clauses[-1].text} {t}"  # 줄바꿈으로 끊긴 문장 이어 붙이기
        else:
            section = t[:80] if len(t) < 60 and not t.endswith(("다.", ".")) else section
            if len(t) >= 60 or t.endswith(("다.", ".")):
                clauses.append(RawClause(t, section, loc))
    return [c for c in clauses if len(c.text) >= 10]  # 너무 짧은 조각은 제외


# ---------- 속성 채우기 ----------


class ClauseAttrs(BaseModel):
    index: int
    event_types: list[EventType] = []
    channels: list[Channel] = []
    regions: list[str] = []
    asset_classes: list[str] = []
    duration: list[Duration] = []
    conditions: list[str] = []
    actions: list[str] = []
    owner_role: str = ""
    protected: Literal["none", "regulatory", "insurance"] = "none"
    protected_evidence: str = ""  # 보호 판단의 원문 인용 (조항 원문의 실제 문구여야 함)


class AttrBatch(BaseModel):
    items: list[ClauseAttrs]


ATTR_SYSTEM = """너는 금융기관 위기 대응 문서의 조항을 구조화한다. 조항 원문은 바꾸지 않는다.
조항마다 다음을 채운다. 원문에 근거가 없으면 비운다 (추측 금지).
- event_types: 이 조항이 전제한 사건 유형 (liquidity_crisis, credit_event, banking_stress, geopolitical, supply_chain,
  trade_policy, pandemic, rate_shock, fx_shock, regulatory, other)
- channels: 영향 경로 (liquidity, credit_spread, counterparty, fx, commodity, supply_chain, rates, equity_vol, sanctions_trade, operational)
- regions: ISO 국가코드 또는 EU·EM·GLOBAL / asset_classes: equity, govt_bond, corporate_bond, money_market 등
- duration: shock, weeks, months, structural / conditions: 조항이 걸어 둔 발동 조건 (짧은 구)
- actions: 해야 할 일 (짧은 명사구, 원문 표현을 살림)
- owner_role: 원문에 적힌 담당 부서명 그대로. 원문에 없으면 빈 문자열
- protected: 법령·감독 규정·공시 의무에 따른 조항이면 regulatory, 보험·예금보험 관련이면 insurance, 아니면 none.
  regulatory·insurance면 protected_evidence에 원문의 근거 문구를 글자 그대로 옮긴다.
index는 입력 순서 그대로."""


def rule_attrs(raw: RawClause) -> ClauseAttrs:
    role = ROLE_RE.search(raw.text)
    prot, ev = "none", ""
    for kind, words in PROTECT_WORDS.items():
        hit = next((w for w in words if w in raw.text), None)
        if hit:
            prot, ev = kind, hit
            break
    return ClauseAttrs(index=0, owner_role=role.group(1) if role else "", protected=prot, protected_evidence=ev)


def llm_attrs(raws: list[RawClause], llm, batch: int = 8) -> list[ClauseAttrs | None]:
    import json

    out: list[ClauseAttrs | None] = [None] * len(raws)
    for start in range(0, len(raws), batch):
        chunk = raws[start: start + batch]
        prompt = json.dumps([{"index": start + i, "section": c.section, "text": c.text} for i, c in enumerate(chunk)],
                            ensure_ascii=False)
        try:
            res = llm.generate(prompt, purpose="doc_clause_attrs", response_schema=AttrBatch, system=ATTR_SYSTEM)
        except Exception:
            continue
        for item in res.items:
            if start <= item.index < start + len(chunk):
                out[item.index] = item
    return out


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", s)


def validate_attrs(raw: RawClause, a: ClauseAttrs | None, roles_in_doc: set[str]) -> tuple[ClauseAttrs, list[str]]:
    """원문 근거 검증. 반환: (검증된 속성, 사람 확인 필요 사유)."""
    fb = rule_attrs(raw)
    issues = []
    if a is None:
        issues.append("LLM 속성 없음 → 규칙 결과만 (전제·조치 비어 있음)")
        a = fb
    a = a.model_copy()
    if a.owner_role and _norm(a.owner_role) not in _norm(raw.text) and a.owner_role not in roles_in_doc:
        issues.append(f"담당 '{a.owner_role}'이 원문에 없음 → 미지정")
        a.owner_role = ""
    if not a.owner_role:
        a.owner_role = fb.owner_role
    if a.protected != "none" and (not a.protected_evidence or _norm(a.protected_evidence) not in _norm(raw.text)):
        if fb.protected != "none":
            a.protected, a.protected_evidence = fb.protected, fb.protected_evidence
            issues.append(f"보호 근거 인용이 원문에 없음 → 규칙 키워드 '{fb.protected_evidence}'로 대체")
        else:
            issues.append("보호 근거 인용이 원문에 없음 → 보호 아님으로")
            a.protected, a.protected_evidence = "none", ""
    if not a.owner_role:
        issues.append("담당 부서 미지정 (사람 확인 필요)")
    return a, issues


@dataclass
class DocIngestResult:
    playbook: Playbook
    rows: list[dict]  # 조항별 원문·위치·속성·확인 사유
    chain: list[dict]


def ingest_document(data: bytes, filename: str, doc_id: str | None = None, llm=None) -> DocIngestResult:
    parse = read_document(data, filename)
    raws = split_clauses(parse)
    doc_id = doc_id or re.sub(r"[^A-Za-z0-9]+", "-", Path(filename).stem).strip("-").upper()[:20] or "DOC"
    doc_id = f"UP-{doc_id}"
    attrs = llm_attrs(raws, llm) if (llm is not None and raws) else [None] * len(raws)
    roles = {m.group(1) for r in raws for m in ROLE_RE.finditer(r.text)}
    clauses, rows, n_issue, n_llm = [], [], 0, 0
    for i, (raw, a) in enumerate(zip(raws, attrs), start=1):
        n_llm += a is not None
        v, issues = validate_attrs(raw, a, roles)
        n_issue += bool(issues)
        prot = v.protected != "none"
        clause = Clause(clause_id=f"{doc_id}-{i:02d}", doc_id=doc_id, text=raw.text,
                        premises=Premises(event_types=v.event_types, channels=v.channels, regions=v.regions,
                                          asset_classes=v.asset_classes, duration=v.duration, conditions=v.conditions),
                        actions=v.actions, owner_role=v.owner_role or "미지정", protected=prot,
                        protected_reason=v.protected if prot else None)
        clauses.append(clause)
        rows.append({"조항": clause.clause_id, "위치": raw.loc, "절": raw.section, "원문": raw.text,
                     "사건 유형": ", ".join(e.value for e in v.event_types), "채널": ", ".join(c.value for c in v.channels),
                     "조치": "; ".join(v.actions), "담당": clause.owner_role,
                     "보호": f"{v.protected} ({v.protected_evidence})" if prot else "-", "확인 필요": "; ".join(issues)})
    pb = Playbook(doc_id=doc_id, title=parse.title, version="upload", year=date.today().year, origin_event="",
                  synthetic=False, source_refs=[filename], clauses=clauses)
    chain = [
        {"단계": "읽기", "원인": f"{filename}: 문단·줄 {len(parse.paragraphs)}개", "근거": "pypdf·python-docx로 본문 추출 (표 포함)",
         "결과": "; ".join(parse.notes) or "완료"},
        {"단계": "조항 분해", "원인": f"조항 {len(raws)}개", "근거": "번호·조문 패턴(제n조, 1., (1), ①, 가., 글머리표)과 문장 이어 붙이기. 원문 그대로 보존",
         "결과": "완료" if raws else "조항을 찾지 못함 (번호 없는 문서는 수동 분할 필요)"},
        {"단계": "속성 채우기", "원인": f"LLM 속성 {n_llm}/{len(raws)}개" if llm is not None else "LLM 미사용 (규칙만)",
         "근거": "담당은 원문 부서명만, 보호(법규·보험)는 원문 인용 근거가 있을 때만, 열거형은 스키마 검증",
         "결과": f"사람 확인 필요 {n_issue}개 조항"},
    ]
    return DocIngestResult(pb, rows, chain)
