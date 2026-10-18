"""대응안 엑셀 내보내기 — 결재·보고에 바로 쓰는 양식. 시트마다 원인 → 근거 → 결과 순서."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import Briefing, Verdict

VERDICT = {Verdict.applicable: "적용", Verdict.partial: "부분적용", Verdict.not_applicable: "부적용", Verdict.uncertain: "불확실"}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def briefing_xlsx(b: Briefing, chain: list[dict] | None = None, clause_text: dict[str, str] | None = None) -> bytes:
    clause_text = clause_text or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    e = b.event
    summary = [("사건", e.title), ("사건일", str(e.date)), ("유형", e.type.value), ("심각도", f"{e.severity}/5"),
               ("지역", ", ".join(e.regions)), ("등급", b.alert_level.value), ("요약", e.summary),
               ("적합 문서 top3", ", ".join(f"{f.doc_id}({f.fit_score:.2f})" for f in b.match.doc_fits[:3])),
               ("공백 채널", ", ".join(g.value for g in b.match.gaps) or "없음"),
               ("비고", b.exposure_note or "-"), ("run_id / trace_id", f"{b.run_id} / {b.trace_id}"),
               ("고지", "에이전트 제안이며 최종 결정은 사람이 한다. 대응문서·포트폴리오는 합성 데이터.")]
    for k, v in summary:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 100
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    if chain:
        _sheet(wb, "판단 근거", ["단계", "원인", "근거", "결과"], [[c["단계"], c["원인"], c["근거"], c["결과"]] for c in chain],
               [18, 50, 70, 30])
    _sheet(wb, "문서 적합도", ["문서", "적합도", "적용", "부분", "부적용", "불확실"],
           [[f.doc_id, round(f.fit_score, 3), f.n_applicable, f.n_partial, f.n_not, f.n_uncertain] for f in b.match.doc_fits],
           [16, 10, 8, 8, 8, 8])
    rows = []
    for f in b.match.doc_fits:
        for j in f.judgments:
            rows.append([f.doc_id, j.clause_id, clause_text.get(j.clause_id, ""), "; ".join(j.mismatch_reasons),
                         j.evidence_quote, VERDICT[j.verdict], round(j.confidence, 2)])
    _sheet(wb, "조항 판정", ["문서", "조항", "조항 원문", "원인(불일치 사유)", "근거(인용)", "결과", "신뢰도"], rows,
           [12, 14, 50, 40, 50, 10, 8])
    exp = [["상위", " → ".join(p.nodes[1:]), p.why, round(p.weight, 6), round(p.relevance, 4), p.uncertainty, round(p.risk_score, 6)]
           for p in b.exposure_top]
    exp += [["동일 발행자", " → ".join(p.nodes[1:]), p.why, round(p.weight, 6), round(p.relevance, 4), p.uncertainty, round(p.risk_score, 6)]
            for p in b.exposure_entities]
    exp += [["사건 지역 핵심", " → ".join(p.nodes[1:]), p.why, round(p.weight, 6), round(p.relevance, 4), p.uncertainty, round(p.risk_score, 6)]
            for p in b.exposure_core]
    _sheet(wb, "노출 경로", ["구분", "경로", "근거", "비중", "연관도", "불확실성", "risk"], exp, [14, 60, 60, 10, 10, 10, 10])
    _sheet(wb, "확인 항목", ["확인", "항목"], [["☐", a] for a in b.actions_to_check], [6, 100])
    if b.evidence_refs:
        _sheet(wb, "출처", ["출처"], [[r] for r in b.evidence_refs], [100])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def holdings_xlsx(df, chain: list[dict]) -> bytes:
    """변환된 공통 스키마 보유내역 + 변환 근거."""
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "보유내역(공통 스키마)", list(df.columns), df.astype(object).where(df.notna(), None).values.tolist(),
           [24, 14, 10, 10, 16, 8, 10, 14, 8, 20, 12, 16, 10])
    _sheet(wb, "변환 근거", ["단계", "원인", "근거", "결과"], [[c["단계"], c["원인"], c["근거"], c["결과"]] for c in chain],
           [14, 50, 80, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
